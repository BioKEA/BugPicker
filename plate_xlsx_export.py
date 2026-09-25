#!/usr/bin/env python3
"""Synchronize a plate CSV into BugPicker's formatted XLSX template."""

from __future__ import annotations

import argparse
import csv
import os
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = SCRIPT_DIR / "Data/plate_spreadsheet_template.xlsx"

HEADER_ALIASES = {
    "quantifiedvolumeul": "quantifiedvolume",
    "volremainingul": "volremaining",
}

NUMERIC_HEADERS = {
    "no",
    "dnaconcentrationngul",
    "extractvolumeul",
    "quantifiedvolume",
    "usedinpcrul",
    "volremaining",
}


def normalize_header(value: Any) -> str:
    normalized = "".join(
        character.lower()
        for character in str(value or "").replace("µ", "u").replace("μ", "u")
        if character.isalnum()
    )
    return HEADER_ALIASES.get(normalized, normalized)


def cell_value(header: str, value: str) -> Any:
    text = str(value or "").strip()
    if not text or normalize_header(header) not in NUMERIC_HEADERS:
        return text or None
    try:
        number = float(text)
    except ValueError:
        return text
    return int(number) if number.is_integer() else number


def synchronize_workbook(csv_path: Path, template_path: Path, output_path: Path) -> dict[str, Any]:
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        csv_headers = list(reader.fieldnames or [])
        rows = list(reader)

    source_path = output_path if output_path.exists() else template_path
    if not source_path.exists():
        raise FileNotFoundError(f"Plate XLSX template not found: {source_path}")

    workbook = load_workbook(source_path)
    worksheet = workbook.active
    csv_by_normalized = {normalize_header(header): header for header in csv_headers}
    mapped_columns: dict[int, str] = {}
    for column in range(1, worksheet.max_column + 1):
        normalized = normalize_header(worksheet.cell(row=1, column=column).value)
        csv_header = csv_by_normalized.get(normalized)
        if normalized and csv_header:
            mapped_columns[column] = csv_header

    if not mapped_columns:
        raise ValueError("No template columns match the plate CSV headers")

    last_row = max(worksheet.max_row, len(rows) + 1)
    for column, csv_header in mapped_columns.items():
        for row_number in range(2, last_row + 1):
            worksheet.cell(row=row_number, column=column).value = None
        for row_number, row in enumerate(rows, start=2):
            worksheet.cell(row=row_number, column=column).value = cell_value(
                csv_header,
                row.get(csv_header, ""),
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(
        prefix=output_path.name + ".",
        suffix=".tmp.xlsx",
        dir=output_path.parent,
    )
    os.close(fd)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        temporary_path.replace(output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()

    return {
        "csv": str(csv_path),
        "xlsx": str(output_path),
        "template": str(template_path),
        "rows": len(rows),
        "mapped_columns": len(mapped_columns),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("csv_path", type=Path)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    csv_path = args.csv_path.resolve()
    output_path = args.output.resolve() if args.output else csv_path.with_suffix(".xlsx")
    result = synchronize_workbook(csv_path, args.template.resolve(), output_path)
    print(
        f"Updated formatted plate workbook: {result['xlsx']} "
        f"({result['rows']} rows, {result['mapped_columns']} mapped columns)",
        flush=True,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
